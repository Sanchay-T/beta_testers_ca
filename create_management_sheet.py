import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Dev Tasks'

# Set column widths
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 38
ws.column_dimensions['C'].width = 40

# Styling
header_fill = PatternFill(start_color='0066CC', end_color='0066CC', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=11)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Title
ws.merge_cells('A1:C1')
ws['A1'] = 'CypherEdge - Development Tasks'
ws['A1'].font = Font(size=13, bold=True, color='0066CC')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 22

# Headers
headers = ['Task', 'What We Need', 'Why']
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col)
    cell.value = header
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border
ws.row_dimensions[2].height = 25

# Task 1
row = 3
ws.cell(row=row, column=1).value = "1. Update System\nS3 + CloudFront"
ws.cell(row=row, column=1).font = Font(size=10, bold=True)

ws.cell(row=row, column=2).value = """Currently using GitHub releases. Need to move to S3 + CloudFront.

S3 stores files, CloudFront distributes them globally. Both work together.

Talked to an industry expert, he suggested this approach."""

ws.cell(row=row, column=3).value = """Updates will be faster and more reliable. Less support tickets.

Costs around $15-20 per month. Takes 3 days to implement.

Your call on whether to go ahead or not."""

for col in range(1, 4):
    cell = ws.cell(row=row, column=col)
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    cell.border = thin_border
    cell.font = Font(size=10)
ws.row_dimensions[row].height = 90

# Task 2
row = 4
ws.cell(row=row, column=1).value = "2. GitHub Actions\nAutomate Builds"
ws.cell(row=row, column=1).font = Font(size=10, bold=True)

ws.cell(row=row, column=2).value = """Automated build system. Tested in CypherSol, works fine.

Need to roll out to other codebases."""

ws.cell(row=row, column=3).value = """Saves manual build time. Reduces code breaking issues.

Suggest implementing across all codebases."""

for col in range(1, 4):
    cell = ws.cell(row=row, column=col)
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    cell.border = thin_border
    cell.font = Font(size=10)
ws.row_dimensions[row].height = 75

# Task 3
row = 5
ws.cell(row=row, column=1).value = "3. CypherX Accuracy\nOutlier Handling"
ws.cell(row=row, column=1).font = Font(size=10, bold=True)

ws.cell(row=row, column=2).value = """CypherX handles most statements fine but some edge cases fail - regional banks, unusual formats.

Need 1000-2000 outlier samples to improve the model."""

ws.cell(row=row, column=3).value = """Will reduce manual intervention and support overhead.

Can handle more banks and expand reach.

Still in discussion. Need your call on priority."""

for col in range(1, 4):
    cell = ws.cell(row=row, column=col)
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    cell.border = thin_border
    cell.font = Font(size=10)
ws.row_dimensions[row].height = 85

# Summary
row = 6
ws.merge_cells(f'A{row}:C{row}')
ws.cell(row=row, column=1).value = """Summary:

Task 2 - Rollout to other codebases
Task 1 - S3 migration, 3 days work
Task 3 - Need priority decision

Need approval on 1 and 2, priority call on 3."""

ws.cell(row=row, column=1).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
ws.cell(row=row, column=1).border = thin_border
ws.cell(row=row, column=1).font = Font(size=10)
ws.row_dimensions[row].height = 75

# Save
wb.save('CypherEdge_Management_Proposal_v2.xlsx')
print('SUCCESS: Excel file created - CypherEdge_Management_Proposal_v2.xlsx')
